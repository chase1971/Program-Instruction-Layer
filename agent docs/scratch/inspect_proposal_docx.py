from pathlib import Path
import hashlib
import json
import sys
import zipfile

from docx import Document


path = Path(sys.argv[1])
doc = Document(path)

print(f"PATH: {path}")
print(f"SIZE: {path.stat().st_size}")
print(f"SHA256: {hashlib.sha256(path.read_bytes()).hexdigest()}")
print(f"SECTIONS: {len(doc.sections)}")
print(f"PARAGRAPHS: {len(doc.paragraphs)}")
print(f"TABLES: {len(doc.tables)}")

print("\nPARAGRAPHS")
for i, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.replace("\t", "[TAB]")
    print(f"P{i:03d} | {paragraph.style.name!r} | {text}")

print("\nTABLES")
for table_index, table in enumerate(doc.tables):
    print(f"TABLE {table_index}: {len(table.rows)}x{len(table.columns)} style={table.style.name if table.style else None!r}")
    for row_index, row in enumerate(table.rows):
        values = [cell.text.replace("\n", " / ") for cell in row.cells]
        print(f"  R{row_index:02d}: {json.dumps(values, ensure_ascii=False)}")

print("\nPACKAGE PARTS")
with zipfile.ZipFile(path) as archive:
    for info in sorted(archive.infolist(), key=lambda item: item.filename):
        print(f"{info.filename}\t{info.file_size}\t{info.CRC}")
