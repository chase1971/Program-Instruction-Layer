from datetime import datetime, timezone
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(sys.argv[1])
NAVY = "1F416B"
PALE_BLUE = "EAF1F8"
LIGHT_GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=115, start=120, bottom=115, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LIGHT_GRAY)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_text, fld_end])
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    if bold_lead and text.startswith(bold_lead):
        run = paragraph.add_run(bold_lead)
        run.bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.34)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_together = True
        paragraph.add_run("•  ")
        if isinstance(item, tuple):
            lead, remainder = item
            run = paragraph.add_run(lead)
            run.bold = True
            paragraph.add_run(remainder)
        else:
            paragraph.add_run(item)


def add_table(doc, headers, rows, widths, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[index]
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, (value, width) in enumerate(zip(row_data, widths)):
            cell = cells[index]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(value)
            run.font.name = "Aptos"
            run.font.size = Pt(font_size)
            run.font.color.rgb = BLACK
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.8)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style.font.size = Pt(24)
title_style.font.bold = True
title_style.font.color.rgb = BLACK
title_style.paragraph_format.space_after = Pt(7)
title_border = title_style.element.get_or_add_pPr().find(qn("w:pBdr"))
if title_border is not None:
    title_style.element.get_or_add_pPr().remove(title_border)

for style_name, size, before, after in (
    ("Heading 1", 15, 14, 6),
    ("Heading 2", 12, 10, 4),
):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = BLACK
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

footer = section.footer
footer.is_linked_to_previous = False
add_page_number(footer.paragraphs[0])

now = datetime.now(timezone.utc)
doc.core_properties.title = "College Algebra Skills Hub Strategic Plan Alignment"
doc.core_properties.subject = "Alignment with the Lone Star College System Strategic Plan 2025 to 2030"
doc.core_properties.creator = ""
doc.core_properties.last_modified_by = ""
doc.core_properties.created = now
doc.core_properties.modified = now

title = doc.add_paragraph(style="Title")
title.add_run("College Algebra Skills Hub Strategic Plan Alignment")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
subtitle_run = subtitle.add_run("Connections to the Lone Star College System Strategic Plan 2025 to 2030")
subtitle_run.italic = True
subtitle_run.font.size = Pt(13.5)
subtitle_run.font.color.rgb = RGBColor(80, 80, 80)

add_heading(doc, "Overall Alignment")
add_body(
    doc,
    "The College Algebra Skills Hub aligns most directly with G2 Leading in Academic Excellence. The project combines instruction with targeted student support, which is the central purpose of G2, and is designed to improve access to explanations, practice, and guided help for students with different levels of preparation. It also supports the System mission by helping students move from their current level of mathematical preparation toward successful completion of College Algebra and later academic goals.",
)
add_body(
    doc,
    "The Hub also makes a meaningful but less direct contribution to G4 Connections to Bridges of Prosperity because College Algebra is an important gateway for many transfer pathways and degree programs. Faculty and tutoring collaboration provide a secondary connection to G3 Culture and Talent Development. G1 Opportunity in Financial Resilience is not a primary alignment; any efficiency created through shared, reusable materials would be an operational benefit rather than evidence that the project fulfills G1.",
)

add_table(
    doc,
    ["Strategic element", "Level of alignment", "Connection to the Hub"],
    [
        ("System vision and mission", "Direct support", "Provides accessible educational support that helps students progress from current preparation toward academic goals"),
        ("G2 Academic Excellence", "Strong and direct", "Combines high-quality instruction with targeted support intended to improve learning and course progress"),
        ("G4 Bridges of Prosperity", "Meaningful contribution", "Strengthens preparation for transfer and degree pathways that depend on successful completion of gateway mathematics"),
        ("G3 Culture and Talent Development", "Secondary support", "Encourages communication, shared development, and instructional innovation among faculty and tutoring staff"),
        ("G1 Financial Resilience", "Limited connection", "May reduce duplicated effort through reusable resources, but does not directly address the employee-centered outcomes of G1"),
    ],
    [1.45, 1.35, 3.85],
)

add_heading(doc, "Strongest Alignment With G2 Academic Excellence")
add_body(
    doc,
    "G2 calls for academic excellence through the combination of high-quality instruction and support that improves learning, completion, transfer, and career outcomes. The Skills Hub is designed around that same combination. It would not replace classroom teaching or assigned coursework. It would extend instruction by giving students a reliable place to review missing prerequisite skills, revisit a current topic, practice independently, or work through a guided explanation.",
)
add_bullets(
    doc,
    [
        ("High-quality instruction: ", "Resources can be aligned with local course expectations, notation, and teaching methods rather than relying entirely on generic materials."),
        ("Layered academic support: ", "Students can choose a short animation, reference material, independent practice, guided practice, or a full tutorial according to the level of help they need."),
        ("Support beyond the classroom: ", "D2L, the Math Portal, and tutoring referrals can connect classroom instruction with support available outside scheduled class time."),
        ("Learning improvement: ", "Practice results, common errors, and student feedback can identify where students are struggling and where resources should be revised."),
        ("Completion and transfer contribution: ", "Stronger support in a gateway course may contribute to persistence and transfer readiness, although those longer-term effects would need to be evaluated rather than assumed."),
    ],
)
add_body(
    doc,
    "This alignment is especially strong because the strategic plan describes student success as a collective responsibility shared by classroom instruction and support services. The Hub would give instructors and tutors a common resource for directing students to targeted help while keeping the instructional connection to the course visible.",
)

add_heading(doc, "Contribution to G4 Bridges of Prosperity")
add_body(
    doc,
    "G4 focuses on connecting students to successful transfer and gainful employment. The Skills Hub would not provide employment placement or transfer advising, so it should not be presented as a direct G4 initiative. Its contribution is earlier in the pathway: helping students build the mathematical foundation needed to progress through College Algebra and into programs that lead to transfer, credentials, or careers.",
)
add_body(
    doc,
    "This is a meaningful connection because the strategic plan emphasizes preparing students to succeed after they leave Lone Star College rather than lowering academic expectations. The Hub supports that approach by maintaining the academic standard while providing more ways for students to reach it. Its strongest G4 relevance would be demonstrated through improved preparation for later coursework and successful movement through transfer-oriented degree pathways.",
)

add_heading(doc, "Secondary Connection to G3 Culture and Talent Development")
add_body(
    doc,
    "G3 emphasizes communication, engagement, and innovation among employees. The Hub could support those conditions by creating a shared process through which faculty and tutoring staff identify recurring student needs, recommend resources, review explanations, and communicate differences in instructional methods. Instructor-specific options could preserve important teaching differences without fragmenting the overall system.",
)
add_body(
    doc,
    "This is a secondary alignment rather than the project's central purpose. The Hub may encourage collaboration and instructional innovation, but it is not itself a comprehensive employee-development or succession-planning program.",
)

add_heading(doc, "Alignment With Broader Strategic Themes")
add_bullets(
    doc,
    [
        ("Removing barriers: ", "The Hub addresses the difficulty students face when they do not know what skill is missing, what the skill is called, or which outside explanation matches their course."),
        ("Connecting instruction and support: ", "The project links faculty instruction, D2L resources, interactive applications, and tutoring rather than treating each as an isolated service."),
        ("Using scale intentionally: ", "A shared collection can make effective resources available across College Algebra sections while still allowing selected instructor-specific alternatives."),
        ("Testing before expanding: ", "The proposed phased approach begins with an inventory, prioritization, and initial build, then uses classroom evidence and feedback to decide what should be revised or developed next."),
        ("Supporting access through technology: ", "D2L provides a consistent institutional home, while the Math Portal could extend selected interactive resources to phones, tablets, and computers."),
    ],
)

add_heading(doc, "Outcomes the Project Could Demonstrate")
add_body(
    doc,
    "The most defensible early outcomes are those directly connected to the Hub's design and use. Broader institutional outcomes should be described as possible contributions until the project has sufficient implementation data.",
)
add_table(
    doc,
    ["Outcome area", "Possible evidence", "Interpretation"],
    [
        ("Access and use", "Visits, starts, completions, repeat use, and referral source", "Shows whether students can find and use the support"),
        ("Learning process", "Correct responses, common errors, help use, and recovery after feedback", "Shows where students struggle and whether guided support assists them"),
        ("Student experience", "Reported clarity, confidence, usefulness, and open-ended feedback", "Shows how students experience the resources and what needs revision"),
        ("Instructional improvement", "Changes made after evidence and results following revision", "Shows whether the Hub is improving through intentional evaluation"),
        ("Course and pathway outcomes", "Course completion, subsequent course progress, or transfer-related indicators", "Potential longer-term contribution requiring careful comparison and sufficient data"),
    ],
    [1.35, 2.55, 2.75],
)

add_heading(doc, "Conclusion")
add_body(
    doc,
    "The College Algebra Skills Hub should be presented primarily as a G2 Academic Excellence project. It directly combines instruction with academic support and creates a structured way to help students address preparation gaps without lowering expectations. Its contribution to G4 is credible through gateway-course success and transfer preparation, and its collaborative model offers secondary support for G3. The strongest strategic case is therefore focused rather than universal: the Hub advances academic excellence, removes barriers to learning, and builds a more connected path between classroom instruction and student support.",
)

source = doc.add_paragraph()
source.paragraph_format.space_before = Pt(10)
source.paragraph_format.space_after = Pt(0)
run = source.add_run("Source: Lone Star College System Strategic Plan 2025 to 2030, final October 2025.")
run.italic = True
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(89, 89, 89)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
