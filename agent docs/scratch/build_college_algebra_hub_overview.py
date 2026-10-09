from datetime import datetime, timezone
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(sys.argv[1])
NAVY = "1F416B"
PALE_BLUE = "EAF1F8"
LIGHT_GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=125, bottom=120, end=125):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LIGHT_GRAY)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_text, fld_end])
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.keep_together = True
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.34)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_together = True
        paragraph.add_run("•  ")
        if isinstance(item, tuple):
            lead, remainder = item
            run = paragraph.add_run(lead)
            run.bold = True
            paragraph.add_run(remainder)
        else:
            paragraph.add_run(item)


def add_table(doc, headers, rows, widths, font_size=10):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[index]
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, (value, width) in enumerate(zip(row_data, widths)):
            cell = cells[index]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(value)
            run.font.name = "Aptos"
            run.font.size = Pt(font_size)
            run.font.color.rgb = BLACK
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(11.5)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style.font.size = Pt(25)
title_style.font.bold = True
title_style.font.color.rgb = BLACK
title_style.paragraph_format.space_after = Pt(7)
title_border = title_style.element.get_or_add_pPr().find(qn("w:pBdr"))
if title_border is not None:
    title_style.element.get_or_add_pPr().remove(title_border)

for style_name, size, before, after in (
    ("Heading 1", 15, 14, 6),
    ("Heading 2", 12, 10, 4),
):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = BLACK
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

footer = section.footer
footer.is_linked_to_previous = False
add_page_number(footer.paragraphs[0])

now = datetime.now(timezone.utc)
doc.core_properties.title = "College Algebra Skills Hub Overview"
doc.core_properties.subject = "Overview of a shared College Algebra support system"
doc.core_properties.creator = ""
doc.core_properties.last_modified_by = ""
doc.core_properties.comments = ""
doc.core_properties.created = now
doc.core_properties.modified = now

title = doc.add_paragraph(style="Title")
title.add_run("College Algebra Skills Hub Overview")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
subtitle_run = subtitle.add_run("A shared resource for review practice and targeted support")
subtitle_run.italic = True
subtitle_run.font.size = Pt(14)
subtitle_run.font.color.rgb = RGBColor(80, 80, 80)

add_heading(doc, "Overview")
add_body(
    doc,
    "The College Algebra Skills Hub would give students one dependable place to review prerequisite algebra, revisit important College Algebra concepts, practice skills, and receive additional instruction when they need it. Students currently have access to many outside resources, but those resources can be difficult to locate and may use methods, notation, or expectations that differ from what is taught in class. The Hub would provide an organized resource that instructors and tutors could recommend with confidence.",
)
add_body(
    doc,
    "D2L would serve as the main home for the Hub. It would contain the complete collection of skill pages, notes, worksheets, videos, animations, practice, tutorials, and links. A companion Math Portal could provide easier access to interactive activities from a phone, tablet, or computer. Together, these tools would support students who need a quick reminder, additional practice, guided help, or a more complete explanation.",
)
add_body(
    doc,
    "This overview describes the intended system without setting a production schedule or defining a final list of skills. Those decisions would be made during the next planning stage after current materials have been reviewed and faculty priorities have been discussed.",
)

add_heading(doc, "Core Parts of the Hub")
add_table(
    doc,
    ["Part", "Purpose"],
    [
        ("D2L Skills Hub", "The organized home for the complete collection and the main entry point for students"),
        ("Instructional resources", "Short animations, notes, examples, worksheets, and longer lessons that explain important skills"),
        ("Independent practice", "Quiz-style practice that allows students to work on their own and receive immediate correctness feedback"),
        ("Guided practice and tutorials", "Interactive activities that help students through a process, respond to errors, and provide more complete instruction"),
        ("Math Portal", "A companion web application for interactive work, progress saving, reporting, surveys, and convenient device access"),
        ("Instructor-specific options", "Selected alternatives for topics where notation, required steps, or teaching methods differ between instructors"),
        ("Tutoring integration", "A shared skill map and direct links that tutors can use to identify missing skills and guide continued practice"),
    ],
    [1.72, 4.93],
)

doc.add_page_break()
add_heading(doc, "Student Experience")
add_body(
    doc,
    "Students would be able to enter the Hub from D2L, a direct instructor link, a tutoring recommendation, or eventually the Math Portal. They could locate a skill by course topic, prerequisite area, or common task and then choose the amount of support they need. A student who has forgotten one rule might use a short animation and worked example. A student who needs practice could complete an independent activity. A student who is struggling with the process could move into guided practice or a full tutorial.",
)
add_body(
    doc,
    "Students would not be expected to complete every resource. The goal is to make the appropriate help easy to find and to give instructors and tutors a clear place to send students for targeted review. When useful, activity data and student feedback could also show where students are struggling and where an explanation or application should be improved.",
)

add_heading(doc, "Content and Organization")
add_body(
    doc,
    "The final content scope has not yet been determined. The Hub would include prerequisite skills that students are expected to know and selected College Algebra concepts that benefit from additional explanation or practice. The next planning stage would create a skill map, review the materials already available in the Boot Camp and current applications, and identify the first topics to develop. The Hub is not intended to reproduce every part of the College Algebra course or every feature of a commercial homework system.",
)
add_body(
    doc,
    "Organization will be essential as the collection grows. Students should be able to browse, search using familiar language, or follow a direct link without needing to know the formal name of a missing skill. Skill pages should use a consistent structure so explanations, examples, practice, guided help, and related skills are easy to recognize. The detailed navigation and naming system would be designed as part of the implementation plan.",
)

add_heading(doc, "Development Approach")
add_body(
    doc,
    "The Hub would be developed in phases rather than built all at once. Early work would focus on organizing what already exists, establishing the skill structure, and creating resources that can be produced and reused efficiently. Short instructional animations and concise reference materials could be expanded first. Independent practice, guided activities, and full tutorials could then be added where the instructional need justifies the additional development time.",
)
add_body(
    doc,
    "Later phases could expand Math Portal access, add instructor-specific versions, and strengthen coordination with tutoring staff and other faculty contributors. Classroom use, student feedback, and tutor or instructor observations would help determine which resources should be revised or developed next. The scope of each phase would be set only after the initial inventory and priority decisions are complete.",
)

doc.add_page_break()
add_heading(doc, "Next Planning Stage")
add_body(
    doc,
    "Before a detailed implementation proposal is created, the following decisions need to be made:",
)
add_bullets(
    doc,
    [
        "Inventory the Boot Camp, existing animations, current applications, and other available materials.",
        "Create the overall skill map and distinguish prerequisite review from concepts taught within College Algebra.",
        "Select the first skills and resource types to develop based on student need, instructional importance, and available materials.",
        "Design the basic D2L organization, naming system, navigation, and connection to the Math Portal.",
        "Determine how instructors and tutoring staff will review, recommend, and use the resources.",
        "Define a realistic first phase and then identify the time, responsibilities, technology, and support needed to complete it.",
    ],
)
add_body(
    doc,
    "The result of that planning work would be a separate implementation scope with clear priorities, phases, responsibilities, and resource needs. This overview can be used first to explain the concept and gather feedback on what the Hub should become.",
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
