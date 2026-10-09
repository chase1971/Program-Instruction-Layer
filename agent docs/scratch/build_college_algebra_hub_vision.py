from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(sys.argv[1])
NAVY = "1F416B"
PALE_BLUE = "EAF1F8"
LIGHT_GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LIGHT_GRAY)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_text, fld_end])
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        run = paragraph.add_run(bold_lead)
        run.bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.34)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run("•  ")
        if isinstance(item, tuple):
            lead, remainder = item
            run = paragraph.add_run(lead)
            run.bold = True
            paragraph.add_run(remainder)
        else:
            paragraph.add_run(item)


def add_numbered(doc, items):
    for number, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.34)
        paragraph.paragraph_format.first_line_indent = Inches(-0.26)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_together = True
        paragraph.add_run(f"{number}.  ")
        paragraph.add_run(item)


def add_table(doc, headers, rows, widths, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[index]
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, (value, width) in enumerate(zip(row_data, widths)):
            cell = cells[index]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(value)
            run.font.name = "Aptos"
            run.font.size = Pt(font_size)
            run.font.color.rgb = BLACK
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style.font.size = Pt(25)
title_style.font.bold = True
title_style.font.color.rgb = BLACK
title_style.paragraph_format.space_after = Pt(7)
title_border = title_style.element.get_or_add_pPr().find(qn("w:pBdr"))
if title_border is not None:
    title_style.element.get_or_add_pPr().remove(title_border)

for style_name, size, before, after in (
    ("Heading 1", 15, 14, 6),
    ("Heading 2", 12, 10, 4),
):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = BLACK
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

footer = section.footer
footer.is_linked_to_previous = False
footer_paragraph = footer.paragraphs[0]
add_page_number(footer_paragraph)

doc.core_properties.title = "College Algebra Skills Hub Vision and Full Scope"
doc.core_properties.subject = "Long-term vision for a shared College Algebra support system"

title = doc.add_paragraph(style="Title")
title.add_run("College Algebra Skills Hub Vision and Full Scope")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
subtitle_run = subtitle.add_run("A shared system for review instruction practice and continuous improvement")
subtitle_run.italic = True
subtitle_run.font.size = Pt(14)
subtitle_run.font.color.rgb = RGBColor(80, 80, 80)

add_heading(doc, "Vision Summary")
add_body(
    doc,
    "The College Algebra Skills Hub is envisioned as a shared, continually improving system that gives students one dependable place to review prerequisite algebra, revisit College Algebra concepts, practice skills, receive guided instruction, and locate help that matches the methods used in their classes. D2L would provide the complete institutional home for the Hub, while the Math Portal would eventually extend interactive and mobile-friendly access to selected resources.",
)
add_body(
    doc,
    "The Hub would serve students, instructors, and the tutoring department. It would combine a common foundation of departmentally useful materials with carefully reviewed instructor-specific alternatives when notation, required steps, or teaching methods differ. Student activity and feedback would be used to improve both the instructional resources and the applications over time.",
)
add_body(
    doc,
    "This document describes the desired end state of the Hub. It does not establish a development schedule, promise a production quantity, or request funding. A later planning document will translate this vision into priorities, phases, responsibilities, timelines, and resource needs.",
)

add_table(
    doc,
    ["Part of the vision", "Intended role"],
    [
        ("D2L Skills Hub", "The complete, organized home for notes, worksheets, videos, animations, practice, tutorials, and links"),
        ("Math Portal", "A companion web application for interactive practice, guided activities, progress saving, reporting, and convenient device access"),
        ("Shared support network", "A common resource used by students, instructors, and tutors, with clear pathways to the right skill"),
        ("Long-term model", "A maintained common foundation with evidence-based revision and selected instructor-specific options"),
    ],
    [1.65, 5.0],
)

add_heading(doc, "Problem and Opportunity")
add_body(
    doc,
    "Students enter College Algebra with widely different preparation. Some need a brief reminder of a prerequisite skill, some need additional practice on a current topic, and others need a concept taught again from the beginning. The difficulty is often not simply a lack of effort: students may not know the name of the missing skill, which resource to trust, or what to do after watching an explanation.",
)
add_body(
    doc,
    "Current options are fragmented. Publisher systems, general websites, videos, and materials from individual courses may use different notation, problem-solving sequences, graphing methods, or written expectations. A mathematically valid explanation can still confuse a student when it does not match the method expected in class. Instructors and tutors also lack one shared destination to recommend for targeted review.",
)
add_body(
    doc,
    "The opportunity is to create an institutional resource that is easier to navigate than an unstructured collection, more aligned with local instruction than a generic commercial system, and more useful than a library of explanations without connected practice or feedback.",
)

add_heading(doc, "Vision Principles")
add_bullets(
    doc,
    [
        ("Easy to find: ", "Students should be able to locate a needed skill even when they do not know its formal mathematical name."),
        ("Layered support: ", "Each skill should offer only as much support as the student needs, from a quick reminder through a full tutorial."),
        ("Instructional consistency: ", "Notation, language, and processes should match classroom expectations, with instructor-specific alternatives when necessary."),
        ("Connected learning: ", "Explanations, examples, practice, feedback, and related prerequisite skills should be linked rather than presented as isolated files."),
        ("Shared but reviewable: ", "Faculty and tutors should be able to contribute ideas and resources through a process that protects accuracy, organization, and consistency."),
        ("Evidence-based improvement: ", "Usage patterns, student work, student feedback, and faculty observations should guide revision."),
        ("Sustainable growth: ", "The system should support phased expansion without becoming too large or disorganized for students to use."),
    ],
)

add_heading(doc, "People the Hub Would Serve")
add_table(
    doc,
    ["Audience", "Primary needs", "How the Hub would help"],
    [
        ("Students", "Review, clarification, practice, and help locating the right prerequisite", "Provides searchable skill pages, multiple levels of support, and consistent instructional methods"),
        ("College Algebra instructors", "A reliable resource to assign, recommend, adapt, and improve", "Provides direct links to skills, common resources, reporting, and selected instructor-specific versions"),
        ("Tutoring staff", "A shared reference for identifying gaps and directing students to continued practice", "Provides a skill map, tutor-facing orientation, and links that tutors can use during and after appointments"),
        ("Program and department faculty", "A maintainable view of common student needs and available resources", "Provides evidence, shared contribution processes, and a basis for coordinated improvement"),
    ],
    [1.35, 2.35, 2.95],
    font_size=8.7,
)

add_heading(doc, "Proposed Student Experience")
add_body(
    doc,
    "Students would enter the Hub from D2L, an instructor link, a tutoring recommendation, a course activity, or eventually the Math Portal. They could browse by course topic, prerequisite skill, or common task; search using familiar language; or follow a direct link to a specific skill page. Each page would use a consistent structure so students know where to find explanations, examples, practice, and additional help.",
)
add_table(
    doc,
    ["Level of support", "Student question", "Resource type"],
    [
        ("Quick reference", "How does this work again", "Concise notes, formulas, worked examples, and short animations"),
        ("Independent practice", "Can I do this on my own", "Practice or quiz-style sets with immediate correctness feedback"),
        ("Guided practice", "Can you help me through the process", "Interactive activities that respond to each step and provide structured recovery after errors"),
        ("Full tutorial", "Can you teach this from the beginning", "Extended instruction that develops the idea, demonstrates the method, and includes supported practice"),
    ],
    [1.35, 2.0, 3.3],
)
add_body(
    doc,
    "A student would not need to complete every resource. Someone who has forgotten one rule might use a short animation and two examples. A student who repeatedly makes the same error might move into guided practice. A student encountering the topic for the first time might use a full tutorial and then independent practice. Instructors and tutors could direct students to the level that best fits the situation.",
)

add_heading(doc, "Full System Scope")

add_heading(doc, "D2L Skills Hub", level=2)
add_body(
    doc,
    "D2L would be the complete institutional home for the system. It would organize skill pages, worksheets, notes, longer videos, animations, practice links, tutorials, and supporting files. Students enrolled in College Algebra would have a predictable entry point that does not depend on locating an outside website or receiving files from one particular section.",
)

add_heading(doc, "Instructional Content Library", level=2)
add_bullets(
    doc,
    [
        "Boot Camp notes, worksheets, and review materials reorganized for use throughout the term.",
        "Short instructional animations that explain one idea, rule, or mathematical transition quickly.",
        "Concise reference pages with definitions, notation, formulas, and worked examples.",
        "Longer video lessons and full written tutorials for concepts requiring sustained explanation.",
        "Downloadable or printable practice and notes where paper-based work is useful.",
    ],
)

add_heading(doc, "Interactive Practice and Tutorials", level=2)
add_body(
    doc,
    "Interactive resources would range from independent skill checks to guided activities that respond to each decision a student makes. Independent practice would provide clear correctness feedback. Guided practice would identify errors, offer targeted help, and allow recovery without simply revealing the answer. Full tutorials would combine explanation, demonstration, supported work, and increasingly independent practice.",
)

add_heading(doc, "Math Portal and Device Access", level=2)
add_body(
    doc,
    "The Math Portal would be a companion web application for interactive resources. It could provide student-specific access, saving and resuming, mobile-friendly layouts, practice history, and direct links back to the complete D2L skill page. D2L would remain the authoritative home for the full collection; the Math Portal would make app-based work easier to access from a phone, tablet, or computer.",
)

add_heading(doc, "Reporting, Feedback, and Continuous Improvement", level=2)
add_body(
    doc,
    "The applications would collect information that helps answer two questions: where are students struggling, and how should the instruction or application be improved? Student-level and aggregate reports could show attempt history, correct responses, common errors, help use, completion, repeated attempts, and recovery after feedback. In-app surveys and comments could show whether directions, explanations, and feedback are clear. Later results could then be reviewed to determine whether revisions improved the experience.",
)

add_heading(doc, "Instructor Specific Customization", level=2)
add_body(
    doc,
    "The shared Hub would provide the common foundation. A later capability could add instructor-specific sections for skills where faculty members use different notation, solution sequences, graphing methods, or required written steps. Students could follow a section-specific link or select their instructor so the explanation and practice match the method used in class. Customization would be used only when an instructional difference matters and would require instructor consultation, review, and maintenance.",
)

add_heading(doc, "Tutoring Department Integration", level=2)
add_body(
    doc,
    "Tutoring staff would be introduced to the Hub, its skill map, and its navigation structure. During an appointment, a tutor could identify a missing prerequisite, open the corresponding resource, and give the student a direct path for continued work after the session. Tutors could also report recurring student difficulties, confusing terminology, or missing resources, making the tutoring department an important source of evidence for future development.",
)

add_heading(doc, "Faculty Contributions and Shared Development", level=2)
add_body(
    doc,
    "Other faculty members could contribute existing materials, recommend skills, identify alternate methods, review mathematical content, or develop new resources. Contributions would enter a shared organizational and review process rather than being added directly as unrelated files. This would preserve the coherence of the Hub while allowing it to benefit from multiple instructors' expertise.",
)

add_heading(doc, "Content Scope")
add_body(
    doc,
    "The Hub would include both prerequisite skills students are expected to know and important concepts taught within College Algebra. Every skill would not require every type of resource. The appropriate level of development would depend on the frequency of the need, the difficulty of the process, the quality of existing materials, and the benefit of interactive support.",
)
add_table(
    doc,
    ["Content area", "Possible skills"],
    [
        ("Algebra foundations", "Signed numbers, order of operations, fractions, exponents, radicals, and algebraic notation"),
        ("Equations and inequalities", "Linear equations, formulas, proportions, compound inequalities, absolute value, and graphing inequalities"),
        ("Lines and coordinate methods", "Slope, forms of a line, intercepts, parallel and perpendicular lines, and graph construction"),
        ("Functions", "Function notation, evaluation, domain and range, composition, inverse functions, and interpretation"),
        ("Transformations and graphing", "Parent functions, shifts, reflections, stretches, compressions, and table-based graphing"),
        ("Polynomials and factoring", "Operations, common factors, trinomial factoring, special products, and solving by factoring"),
        ("Quadratic equations", "Factoring, square-root methods, completing the square, the quadratic formula, and interpretation of solutions"),
        ("Rational and radical expressions", "Simplifying, restrictions, operations, equations, and extraneous solutions"),
        ("Exponential and logarithmic functions", "Properties, equations, graphs, and applications where included in the course"),
        ("Systems and matrices", "Systems of equations, elimination, augmented matrices, and Gaussian or Gauss-Jordan elimination where included"),
    ],
    [2.05, 4.6],
)

add_heading(doc, "Organization, Search, and Navigation")
add_body(
    doc,
    "Organization is a central requirement, not a cosmetic detail. A large collection will fail if students cannot recognize where they are, locate the correct skill, or understand which resource to use. The Hub should therefore use a documented information structure before large-scale production begins.",
)
add_bullets(
    doc,
    [
        ("Consistent skill pages: ", "Use the same order and labels for overview, prerequisite links, explanation, examples, practice, guided help, and related skills."),
        ("Multiple ways to find a skill: ", "Support browsing by course unit, prerequisite area, common task, or resource type, along with direct instructor and tutor links."),
        ("Student-friendly search language: ", "Include synonyms and common phrases such as finding slope, clearing fractions, or moving a graph, not only formal topic names."),
        ("Visible relationships: ", "Show prerequisite skills, related skills, and suggested next steps so students can move through the system intentionally."),
        ("Useful filters and labels: ", "Identify the support level, estimated time, device suitability, prerequisite status, and any instructor-specific version."),
        ("Limited duplication: ", "Reuse shared resources and link to them instead of creating conflicting copies in several modules."),
    ],
)

add_heading(doc, "Collaboration, Contribution, and Governance")
add_body(
    doc,
    "A shared Hub needs a clear way for people to participate without allowing the system to become inconsistent. Faculty and tutoring staff should be able to suggest content and identify needs, while a defined review process confirms mathematical accuracy, accessibility, instructional fit, naming, metadata, and placement before publication.",
)
add_bullets(
    doc,
    [
        "Identify an owner or coordinating group responsible for organization, publishing standards, and long-term maintenance.",
        "Create a simple contribution process for existing faculty materials, new resource proposals, corrections, and alternate instructional methods.",
        "Separate the common departmental foundation from optional instructor-specific additions.",
        "Record the author, reviewer, revision date, connected skills, and intended use of each resource.",
        "Review resources periodically so outdated, duplicated, or unused materials do not accumulate indefinitely.",
    ],
)

add_heading(doc, "Tutoring Support")
add_body(
    doc,
    "The tutoring department would be treated as an active partner rather than simply another audience. Tutors could receive a short orientation, a searchable skill map, and quick-reference guidance for directing students to resources. The department could include Hub links in appointment follow-up, tutoring webpages, handouts, or QR materials. Tutor observations could help identify recurring prerequisite gaps and places where students cannot find or understand existing help.",
)

add_heading(doc, "Student Awareness and Adoption")
add_body(
    doc,
    "Automatic D2L access alone will not ensure that students understand the Hub or use it. The system needs a deliberate awareness and onboarding plan that presents the Hub as a normal part of College Algebra support rather than an optional site students discover only after they are already struggling.",
)
add_bullets(
    doc,
    [
        "Introduce the Hub near the beginning of College Algebra courses with one short required or guided orientation activity.",
        "Provide faculty with standard syllabus language, announcement text, and direct links to frequently used skills.",
        "Ask instructors to incorporate selected Hub activities into course expectations when they support current instruction.",
        "Promote the Hub through tutoring staff, tutoring webpages, advising, student-success offices, and relevant campus communications.",
        "Use clear visual identity and consistent naming so students recognize the same resource across D2L, the Math Portal, and tutoring materials.",
        "Track how students first reach the Hub and which referral methods lead to meaningful use." ,
    ],
)

add_heading(doc, "Existing Foundation")
add_body(
    doc,
    "The vision builds on an existing foundation rather than beginning with an empty course. Current materials and prototypes demonstrate several parts of the intended system and can be evaluated for reuse, revision, or migration.",
)
add_bullets(
    doc,
    [
        "Boot Camp notes, worksheets, and review materials that can be inventoried and reorganized.",
        "Short instructional animations covering algebraic manipulation, equations, inequalities, fractions, factoring, functions, graphing, and related topics.",
        "Reference materials and worked examples for important College Algebra skills.",
        "Interactive applications for factoring, fractions, quadratic equations, transformations, and matrices.",
        "Guided and independent Gauss-Jordan practice and guided transformation work.",
        "Math Portal capabilities for student access, saving and resuming, feedback, reporting, surveys, and activity review.",
        "Teacher-facing reports that can reveal individual and aggregate response patterns." ,
    ],
)

add_heading(doc, "Evidence and Continuous Improvement")
add_body(
    doc,
    "Evidence would serve two purposes: assessing whether students find the resources useful and improving the resources themselves. Existing and future classroom use of the applications can provide evidence of access, completion, response patterns, confidence, perceived usefulness, and common difficulties. Results should be reported with sample sizes and appropriate limits rather than used to make unsupported claims.",
)
add_table(
    doc,
    ["Evidence area", "What it could show"],
    [
        ("Reach and adoption", "How students find the Hub, which referral routes work, and who opens assigned or optional resources"),
        ("Engagement", "Starts, completions, repeat visits, saved work, video viewing, and return after feedback"),
        ("Learning process", "Correct responses, common first-attempt errors, help use, recovery after errors, and unfinished steps"),
        ("Student experience", "Reported understanding, confidence, usefulness, clarity, and open-ended comments"),
        ("Instructional improvement", "Which explanations, directions, examples, or app behaviors should be revised and whether later revisions help"),
        ("Faculty and tutor observations", "Repeated skill gaps, navigation problems, missing topics, and differences between available resources and classroom expectations"),
    ],
    [1.7, 4.95],
)
add_body(
    doc,
    "Classroom evidence already being collected from current applications, including matrix and transformation activities, can be incorporated into the planning process when the questions, sample sizes, and conditions of use are documented. Where another system is used for comparison, students should complete comparable work in both formats before preference or effectiveness claims are made.",
)

add_heading(doc, "Quality, Accessibility, and Stewardship")
add_bullets(
    doc,
    [
        "Mathematical content should be reviewed for accuracy and alignment with the intended instructional method.",
        "Resources should follow institutional accessibility requirements and be usable with common devices and assistive technologies.",
        "Student-identifying information should be handled through approved institutional systems and excluded from unnecessary cloud activity records.",
        "Shared templates should establish consistent notation, visual design, feedback patterns, and navigation.",
        "Every resource should have a maintenance path so corrections and updates can be made without rebuilding the entire system.",
        "The system should distinguish experimental materials from reviewed materials that are ready for broad student use.",
    ],
)

add_heading(doc, "Long Term End State")
add_body(
    doc,
    "In its mature form, the College Algebra Skills Hub would be a dependable part of the institution's mathematics support system. A student, instructor, or tutor could identify a skill, locate an explanation that matches classroom expectations, choose an appropriate level of practice, and see what to do next. Faculty could contribute and improve resources without fragmenting the system. Evidence from actual use would guide priorities and revisions. D2L, the Math Portal, classroom instruction, and tutoring support would function as connected parts of one organized experience.",
)

add_heading(doc, "Next Planning Phase")
add_body(
    doc,
    "The next step is to convert this full vision into an implementation scope. That planning process should answer the following questions before committing to a production schedule or promising specific quantities.",
)
add_numbered(
    doc,
    [
        "What materials already exist in the Boot Camp, current courses, animations, applications, and faculty collections, and which are ready to reuse, revise, replace, or retire?",
        "What complete skill map should organize the Hub, and how should it distinguish prerequisite skills from concepts taught within College Algebra?",
        "Which skills should be developed first based on frequency, instructional importance, available resources, and the benefit of guided support?",
        "What information architecture, naming system, search vocabulary, tags, and navigation paths will keep the Hub usable as it grows?",
        "What should the first functional version contain, and what criteria will determine that it is ready for student use?",
        "Which resources belong in D2L, which belong in the Math Portal, and how should the two systems link to one another?",
        "How will faculty members and tutoring staff contribute materials, request resources, review content, and report problems?",
        "Which instructional differences justify instructor-specific versions, and how will students reach the correct version?",
        "What evidence from current classroom use should be compiled, and what measures should be collected during future pilots?",
        "How will students learn that the Hub exists, understand its purpose, and encounter it often enough for it to become a normal support tool?",
        "What accessibility, privacy, D2L, ownership, and long-term maintenance requirements must be resolved with institutional partners?",
        "What staffing, collaboration, development time, technology, and funding options would support each implementation phase?",
    ],
)
add_body(
    doc,
    "The output of the next planning phase should be a prioritized inventory, an agreed skill map, an initial information architecture, a defined first build, a contribution and review process, an evaluation plan, and a phased implementation roadmap. Those decisions can then support later proposals for time, staffing, technology, or funding without narrowing the long-term vision described here.",
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
