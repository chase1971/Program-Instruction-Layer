from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"C:\Users\chase\Documents\Programs\School Scrips\School documents\College Algebra Skills Hub Proposal Draft.docx")
NAVY = "17365D"
PALE_BLUE = "EAF1F8"
PALE_GRAY = "F5F6F7"
LIGHT_GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LIGHT_GRAY)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_text, fld_end])
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[index]
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, (value, width) in enumerate(zip(row_data, widths)):
            cell = cells[index]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(value)
            run.font.name = "Aptos"
            run.font.size = Pt(9)
            run.font.color.rgb = BLACK
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.34 if level == 0 else 0.62)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(3)
        p.add_run("•  ")
        p.add_run(item)


def add_numbered(doc, items):
    for number, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.34)
        p.paragraph_format.first_line_indent = Inches(-0.26)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{number}.  ")
        p.add_run(item)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    keep_with_next(p)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style.font.size = Pt(25)
title_style.font.bold = True
title_style.font.color.rgb = BLACK
title_style.paragraph_format.space_after = Pt(7)
title_p_pr = title_style.element.get_or_add_pPr()
title_border = title_p_pr.find(qn("w:pBdr"))
if title_border is not None:
    title_p_pr.remove(title_border)

for style_name, size, before, after in (
    ("Heading 1", 15, 15, 6),
    ("Heading 2", 12, 10, 4),
    ("Heading 3", 10.5, 8, 3),
):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = BLACK
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for list_style in ("List Bullet", "List Bullet 2", "List Number"):
    style = styles[list_style]
    style.font.name = "Aptos"
    style.font.size = Pt(10.5)
    style.font.color.rgb = BLACK
    style.paragraph_format.space_after = Pt(3)

footer = section.footer
footer.paragraphs[0].text = ""
add_page_number(footer.paragraphs[0])

title = doc.add_paragraph(style="Title")
title.add_run("Proposal for Development of the College Algebra Skills Hub")
title_p_pr = title._p.get_or_add_pPr()
title_border = title_p_pr.find(qn("w:pBdr"))
if title_border is not None:
    title_p_pr.remove(title_border)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(13)
run = subtitle.add_run("A D2L based resource for review instruction and practice")
run.font.name = "Aptos"
run.font.size = Pt(13)
run.font.italic = True
run.font.color.rgb = RGBColor(64, 64, 64)

summary_table = doc.add_table(rows=3, cols=2)
summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
summary_table.autofit = False
set_table_borders(summary_table)
summary_rows = [
    ("Primary request", "One course of summer release time for concentrated development"),
    ("Additional request", "Funding for professional AI development tools at $200 per month during the approved project period"),
    ("Project period", "Development beginning in January, concentrated work during the summer release, and continued revision during the following academic year"),
]
for row_index, (label, value) in enumerate(summary_rows):
    left, right = summary_table.rows[row_index].cells
    left.width = Inches(1.55)
    right.width = Inches(5.05)
    for cell in (left, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=120, bottom=120)
    set_cell_shading(left, NAVY)
    if row_index % 2 == 1:
        set_cell_shading(right, PALE_BLUE)
    left_p = left.paragraphs[0]
    left_p.paragraph_format.space_after = Pt(0)
    left_r = left_p.add_run(label)
    left_r.bold = True
    left_r.font.name = "Aptos"
    left_r.font.size = Pt(9)
    left_r.font.color.rgb = RGBColor(255, 255, 255)
    right_p = right.paragraphs[0]
    right_p.paragraph_format.space_after = Pt(0)
    right_r = right_p.add_run(value)
    right_r.font.name = "Aptos"
    right_r.font.size = Pt(9.5)
summary_table.rows[0].cells[0]._tc.getparent()

add_heading(doc, "Executive Summary")
doc.add_paragraph(
    "I request one course of summer release time to develop the College Algebra Skills Hub, "
    "a D2L course that would be automatically available to students enrolled in College Algebra. "
    "The Hub would give students one organized location for reviewing prerequisite algebra skills, "
    "revisiting important course concepts, watching short explanations, and practicing at a level of support that matches their needs."
)
doc.add_paragraph(
    "The proposed work builds on materials and systems that already exist. Current resources include Boot Camp notes and review materials, "
    "short mathematical animations, reference pages, independent practice, guided practice, full interactive tutorials, student progress saving, "
    "and instructor reports. The release time would allow these pieces to be organized into a coherent course, expanded across the highest priority "
    "College Algebra skills, tested with students, and prepared for continued use by multiple instructors."
)
doc.add_paragraph(
    "Development would begin in January and continue throughout the academic year. The summer release would provide the uninterrupted time needed "
    "for the most concentrated production work. Because a short animation can be produced much faster than a guided activity or full tutorial, the proposal "
    "defines a firm sequence of deliverables without promising an arbitrary number of items before the production process has been benchmarked. "
    "The first priority will be a broad collection of short explanations and reference materials. Independent practice will follow, with a smaller number "
    "of guided activities and full tutorials developed for skills where students benefit most from step-by-step support."
)

add_heading(doc, "Need for the Project")
doc.add_paragraph(
    "Students enter College Algebra with different levels of preparation. Some have recently completed prerequisite coursework, while others are returning "
    "to mathematics after a long gap. A student may understand the current topic but be missing one earlier skill, such as operations with fractions, solving "
    "equations, factoring, reading function notation, or graphing a line. Other students need another explanation of a concept taught in College Algebra before "
    "they are ready to practice it independently."
)
doc.add_paragraph(
    "At present, students can search for help on their own, use a publisher system, ask an instructor, or consult materials distributed in an individual course. "
    "These options do not provide a consistent path. Students often do not know the name of the skill they need, which source matches the method used in class, "
    "or what they should do after watching an explanation. Instructors also lack one common destination they can recommend when a student needs targeted review."
)
doc.add_paragraph(
    "A D2L course assigned to every College Algebra student would address this access problem. An instructor could direct a student to a named skill in the Hub, "
    "and the student could move from a brief reminder to practice or guided instruction without having to locate and evaluate outside resources. The Hub would also "
    "make existing Boot Camp materials available throughout the term instead of limiting them to a single review period."
)

add_heading(doc, "Project Purpose and Goals")
doc.add_paragraph(
    "The purpose of the College Algebra Skills Hub is to create a sustainable instructional resource that supports students before and during College Algebra. "
    "The project will organize existing materials, produce new resources, and establish a repeatable method for adding content over time."
)
add_bullets(doc, [
    "Give every College Algebra student a consistent location for prerequisite review and course support.",
    "Present algebra skills using the same notation, language, and problem-solving methods students encounter in class.",
    "Offer different levels of support so students can choose a quick reminder, independent practice, guided practice, or a full tutorial.",
    "Give instructors a specific resource to assign or recommend when a student has an identifiable skill gap.",
    "Use student activity, performance patterns, and feedback to guide revision and future development.",
    "Create materials that can be reused across sections and expanded in later semesters or summer development periods.",
])

add_heading(doc, "Proposed Student Experience")
doc.add_paragraph(
    "The Hub will be organized by recognizable skills rather than by software product. Each skill page may include several kinds of resources, depending on the "
    "importance and complexity of the topic. Students will not be required to complete every resource. Instructors can assign selected activities, while students "
    "can also return to the Hub when they need additional review."
)
add_table(
    doc,
    ["Level of support", "Student question", "Resource type"],
    [
        ("Quick reference", "How does this work again", "Concise notes, formulas, worked examples, and short animations"),
        ("Independent practice", "Can I do this on my own", "Practice problems or quiz style sets with immediate correctness feedback"),
        ("Guided practice", "Can you help me through the process", "Interactive activities that respond to each step and provide structured recovery after an error"),
        ("Full tutorial", "Can you teach this from the beginning", "Extended instruction that develops the idea, demonstrates the method, and includes supported practice"),
    ],
    [1.35, 2.0, 3.25],
)
doc.add_paragraph(
    "This layered structure is important because students do not always need the same intervention. A student who has forgotten one rule may need a two-minute "
    "animation. A student who repeatedly makes the same error may need guided practice. A student who never learned the topic may need a full tutorial. The Hub "
    "will allow these resources to coexist without requiring every skill to begin as a large application."
)

add_heading(doc, "Existing Foundation")
doc.add_paragraph(
    "The project begins with a working foundation rather than a blank course shell. Existing materials demonstrate both the range of possible resources and the "
    "technical ability to deliver them to students."
)

add_heading(doc, "Existing Instructional Content", level=2)
add_bullets(doc, [
    "Boot Camp notes, worksheets, and review materials that can be reorganized and transferred into the D2L course.",
    "Short instructional animations on algebraic manipulation, equations, inequalities, fractions, factoring, functions, domain and range, and related topics.",
    "Reference resources for equations of lines, function families, domain and range, and worked examples.",
    "Interactive applications for factoring, fractions, solving quadratics, transformations, logic, and matrices.",
    "Guided and independent practice for the Gauss-Jordan method, including a reusable problem bank.",
    "Guided homework for identifying and graphing transformations and for completing logic truth tables.",
])

add_heading(doc, "Existing Delivery and Reporting Tools", level=2)
add_bullets(doc, [
    "A student portal that can deliver videos, reference pages, quizzes, guided activities, and tutorials.",
    "Student-specific access and the ability to save and resume work.",
    "Immediate feedback and guided recovery within selected activities.",
    "Student attempt history and instructor review of individual responses.",
    "Aggregate reports showing patterns of correct responses, errors, help use, and completion.",
    "Surveys and viewing data that can be used to improve the materials.",
    "A privacy-conscious design in which student names are not stored in the cloud activity database.",
])
doc.add_paragraph(
    "Some current examples were built for topics outside College Algebra. Those examples are not proposed as College Algebra content, but they show that the same "
    "production process can support multiple types of mathematical explanation. The College Algebra project will apply that process to a carefully selected algebra curriculum."
)

add_heading(doc, "Preliminary Student Evidence")
doc.add_paragraph(
    "The current evidence is preliminary and comes from early classroom use rather than a controlled comparison. Initial responses to the transformation activities are "
    "encouraging. In one early set of five responses, all five students reported improved understanding after using the activity. Three reported that they preferred the "
    "app and two reported that they preferred MyLab Math. Three described themselves as somewhat confident afterward and two as very confident."
)
doc.add_paragraph(
    "The preference question should not be treated as a valid comparison because students did not complete matched versions of the same work in both systems. The MyLab "
    "graphing activity also uses a different task design and does not require students to construct graphs by hand using the table method taught in class. The current result "
    "therefore indicates student reaction, not comparative effectiveness."
)
doc.add_paragraph(
    "Participation has also been limited because the applications were introduced sporadically and were not consistently built into the syllabus or course grade. In the current "
    "College Algebra group, only a portion of the approximately 26 students had begun the optional activity at the time of this draft. This is an implementation limitation rather "
    "than evidence that the resource is ineffective. A more structured spring implementation will introduce the activities earlier, make selected use an expected part of the course, "
    "and provide better information about completion, repeated use, and student response."
)
doc.add_paragraph(
    "The final proposal can be updated with results from the matrix tutorial, transformation activities, spring implementation, student comments, completion data, and any available "
    "before-and-after performance measures. These findings will be reported with sample sizes and clear limits."
)

add_heading(doc, "Project Scope")
doc.add_paragraph(
    "The project will develop the course in phases. The scope is defined by completed systems and resource categories rather than a fixed count of animations or applications. "
    "This approach protects the quality of the work while still creating clear expectations for what the release time will produce."
)

work_before_heading = add_heading(doc, "Work Before the Summer Release", level=2)
work_before_heading.paragraph_format.page_break_before = True
work_before_heading.paragraph_format.keep_with_next = False
add_bullets(doc, [
    "Inventory the existing Boot Camp materials, current applications, animations, notes, worksheets, and practice resources.",
    "Create a prioritized College Algebra skill map that distinguishes prerequisite review from concepts taught within the course.",
    "Identify which skills need only a reference or short animation and which justify independent practice, guided practice, or a full tutorial.",
    "Establish a standard D2L module structure and a consistent visual and instructional format.",
    "Continue collecting classroom use data and student feedback during the spring term.",
    "Benchmark the production time for representative short, medium, and intensive resources so the summer plan can be adjusted using actual development data.",
])

add_heading(doc, "Work During the Summer Release", level=2)
add_numbered(doc, [
    "Build the D2L course structure and organize the priority skills into a clear sequence that students and instructors can navigate.",
    "Transfer and revise applicable Boot Camp notes, worksheets, and review materials for year-round use.",
    "Produce the first broad collection of short instructional animations and concise reference materials, with emphasis on high-frequency prerequisite gaps and important College Algebra skills.",
    "Develop independent practice for selected skills where immediate right-or-wrong feedback is instructionally useful.",
    "Develop a smaller set of guided-practice activities for skills that require students to make and connect several decisions.",
    "Complete or refine selected full tutorials that can serve as models for later development, such as factoring or Gaussian elimination.",
    "Test navigation, readability, accessibility, mathematical accuracy, saving and resuming, and instructor reporting.",
    "Document the development process, remaining priorities, and recommended next phase.",
])

add_heading(doc, "Work After the Summer Release", level=2)
add_bullets(doc, [
    "Use the Hub in College Algebra sections and gather structured information about access, completion, repeat use, common errors, and student perceptions.",
    "Revise resources based on observed difficulties and instructor feedback.",
    "Continue adding animations and practice during the academic year as time permits.",
    "Prepare future requests for additional development time only after the first phase establishes production rates and instructional priorities.",
])

add_heading(doc, "Scope Priorities and Limits")
doc.add_paragraph(
    "The first development priority will be breadth through short animations, notes, and worked examples. These resources can address many common needs and can be produced "
    "more efficiently than full interactive applications. Independent practice will be added next. Guided practice and full tutorials will be reserved for concepts where their "
    "additional development time produces a clear instructional benefit."
)
doc.add_paragraph(
    "The project will not attempt to reproduce every topic in a College Algebra textbook during one summer. It is also not intended to replace the instructor, the primary course, "
    "or every function of a publisher homework system. The Hub will provide targeted instruction and practice that matches the department's methods and fills gaps that existing systems "
    "do not address well."
)
doc.add_paragraph(
    "Final production quantities will be established after the early benchmark. A simple animation, an animation with several coordinated mathematical transitions, an independent "
    "practice bank, and a guided tutorial require very different amounts of design, programming, testing, and revision. Reporting only a total item count would obscure those differences. "
    "The final project report will instead document completed resources by type, instructional purpose, and level of student support."
)

add_heading(doc, "Planned Deliverables")
add_table(
    doc,
    ["Deliverable", "Completion standard"],
    [
        ("D2L course framework", "A navigable course organized by priority algebra skills and ready for student access"),
        ("College Algebra skill map", "A documented sequence of prerequisite and course-level skills with development priorities"),
        ("Standard module design", "A repeatable structure for notes, animations, practice, guided work, and related files"),
        ("Boot Camp migration", "Applicable existing notes and worksheets revised and placed in the new structure"),
        ("Animation and reference collection", "A substantial first collection focused on the highest-priority skills, with each item reviewed for mathematical accuracy and readability"),
        ("Independent practice", "Selected skill-based practice with clear feedback and instructor-visible results where appropriate"),
        ("Guided instruction", "Selected guided activities or tutorials for processes that benefit from step-by-step support"),
        ("Evaluation process", "A consistent plan for collecting use, completion, performance, and student feedback information"),
        ("Continuation plan", "An end-of-phase report identifying completed work, lessons learned, remaining priorities, and recommended next steps"),
    ],
    [2.05, 4.55],
)

add_heading(doc, "Candidate Content Areas")
doc.add_paragraph(
    "The content map will be finalized after the Boot Camp inventory and consultation with College Algebra faculty. Candidate areas include the following. Inclusion in this list "
    "does not mean that every area will receive every level of support during the first phase."
)
add_table(
    doc,
    ["Content area", "Possible skills"],
    [
        ("Algebra foundations", "Signed numbers, order of operations, fractions, exponents, radicals, and algebraic notation"),
        ("Equations and inequalities", "Linear equations, formulas, proportions, compound inequalities, absolute value, and graphing inequalities"),
        ("Lines and coordinate methods", "Slope, forms of a line, intercepts, parallel and perpendicular lines, and graph construction"),
        ("Functions", "Function notation, evaluation, domain and range, composition, inverse functions, and interpretation"),
        ("Transformations and graphing", "Parent functions, shifts, reflections, stretches, compressions, and table-based graphing"),
        ("Polynomials and factoring", "Operations, common factors, trinomial factoring, special products, and solving by factoring"),
        ("Quadratic equations", "Factoring, square root methods, completing the square, the quadratic formula, and interpretation of solutions"),
        ("Rational and radical expressions", "Simplifying, restrictions, operations, equations, and extraneous solutions"),
        ("Exponential and logarithmic functions", "Properties, equations, graphs, and applications where included in the course"),
        ("Systems and matrices", "Systems of equations, elimination, augmented matrices, and Gaussian or Gauss-Jordan elimination where included in the course"),
    ],
    [2.05, 4.55],
)

add_heading(doc, "Development Method")
doc.add_paragraph(
    "Each resource will begin with a defined student need and an observable learning task. Existing materials and reusable software patterns will be used when they fit the instructional "
    "goal. New work will be tested first on a representative example before similar resources are produced in a batch. Mathematical accuracy, instructional sequence, readability, and device "
    "layout will be checked before publication."
)
doc.add_paragraph(
    "The project will use a common structure so later resources do not become isolated products. Short animations will follow a shared visual and notation style. Practice activities will use "
    "common patterns for feedback, saving, resuming, and reporting. Guided activities will record where students answered correctly, where they needed help, and where instructions may need revision."
)
doc.add_paragraph(
    "Instructor review remains part of every stage. AI tools may accelerate drafting, coding, animation production, testing, and revision, but mathematical decisions, instructional methods, and final "
    "approval will remain under instructor control. Student-identifying information will not be submitted to an AI system as part of the development process."
)

add_heading(doc, "Evaluation Plan")
doc.add_paragraph(
    "The initial evaluation will focus on whether students use the resource, whether they can complete the activities, what errors occur, and whether students report that the materials improve "
    "their understanding and confidence. The available class sizes may not support strong causal claims, so results will be used primarily for instructional improvement and feasibility assessment."
)
add_table(
    doc,
    ["Evaluation area", "Evidence"],
    [
        ("Reach", "Number and proportion of students who open assigned and optional resources"),
        ("Engagement", "Starts, completions, repeat visits, saved work, and video viewing where available"),
        ("Learning process", "Correct responses, first-attempt errors, help use, recovery after errors, and completion of guided steps"),
        ("Student perception", "Reported understanding, confidence, usefulness, clarity, and open-ended comments"),
        ("Course integration", "Completion when activities are optional compared with completion when selected activities are expected or graded for participation"),
        ("Instructor use", "Skills assigned or recommended, common reasons for referral, and faculty suggestions for additional content"),
    ],
    [1.65, 4.95],
)
doc.add_paragraph(
    "Where a comparison is made between the Hub and another system, students should complete comparable work in both formats before being asked which they prefer. Any reported comparison will identify "
    "the task, sample size, exposure to each format, and limits of interpretation. This will avoid treating a preference question as evidence of effectiveness when the experiences were not equivalent."
)

add_heading(doc, "Sustainability and Broader Use")
doc.add_paragraph(
    "D2L will provide a stable entry point that students and instructors already know. Organizing the materials by skill will make the course useful across different sections and textbook schedules. "
    "The modular design will also allow an instructor to link directly to a specific skill without assigning the entire Hub."
)
doc.add_paragraph(
    "The project is designed to continue beyond one summer. Short resources can be added throughout the year, while later release time or development periods can focus on more intensive guided practice "
    "and tutorials. Student data and faculty requests will determine which topics receive deeper development. This prevents the project from depending on a one-time attempt to build a comprehensive course all at once."
)

add_heading(doc, "Requested Support")

add_heading(doc, "Summer Course Release", level=2)
doc.add_paragraph(
    "I request a one-course summer release so that I may teach one summer course while receiving the compensation normally associated with two. The released time will be used for the concentrated development, "
    "testing, organization, and documentation described in this proposal. Work will begin before the summer, but the release is necessary because the guided activities, tutorials, and course-wide organization require "
    "sustained blocks of time that are difficult to maintain alongside a full summer teaching load."
)

add_heading(doc, "AI Development Funding", level=2)
doc.add_paragraph(
    "I also request funding for professional AI development tools at a cost of $200 per month during the approved project period. These tools support software development, animation production, test creation, "
    "revision, documentation, and quality checks. The requested amount should be calculated as $200 multiplied by the number of funded months. The final budget can therefore match the approved start and end dates "
    "without assuming a funding period in advance."
)

add_heading(doc, "D2L and Institutional Support", level=2)
doc.add_paragraph(
    "The project will also require assistance with creation of the D2L course, automatic assignment or enrollment of College Algebra students, appropriate instructor access, and confirmation of institutional requirements "
    "for accessibility, student data, and long-term course ownership. These needs primarily involve coordination with existing institutional systems rather than a separate software purchase."
)

add_heading(doc, "Risks and Responses")
add_table(
    doc,
    ["Risk", "Response"],
    [
        ("The topic list becomes too large", "Use the skill map to rank needs and complete the highest-value resources first"),
        ("Production rates vary by resource type", "Benchmark representative resources and report output by type rather than as one item count"),
        ("Students do not use optional materials", "Introduce selected activities early and make some use an explicit course expectation or participation assignment"),
        ("Early samples are small", "Use results for revision, report sample sizes, combine evidence across terms when appropriate, and avoid causal claims"),
        ("Resources become inconsistent", "Use shared templates, notation rules, navigation patterns, and review criteria"),
        ("The project depends on one summer", "Design the Hub for phased expansion and document the next priorities at the end of the release period"),
    ],
    [2.25, 4.35],
)

add_heading(doc, "Approval Requested")
doc.add_paragraph(
    "Approval is requested for one course of summer release time, funding for professional AI development tools at $200 per month during the approved project period, and institutional support for the D2L course and enrollment process. "
    "The release will allow the existing prototype, Boot Camp materials, and early classroom experience to be developed into an organized College Algebra resource that can be used across sections and expanded over time."
)

add_heading(doc, "Information to Add Before Final Submission")
doc.add_paragraph(
    "Before submission, this proposal should be updated with the institution's required dates and format, the final funding period, the selected College Algebra skill priorities, and the latest classroom evidence. The evidence section should include:"
)
add_bullets(doc, [
    "Use and completion data from the matrix tutorial and transformation activities.",
    "Student feedback with sample sizes and the exact questions asked.",
    "Spring results after the applications are introduced earlier and used more consistently.",
    "Any available evidence of improvement between the beginning and end of an activity.",
    "Representative student comments that illustrate what students found helpful or difficult.",
    "Faculty feedback on the proposed skill map and the usefulness of a shared D2L resource.",
    "A final project calendar based on the institution's summer dates.",
])

# Document core properties
doc.core_properties.title = "Proposal for Development of the College Algebra Skills Hub"
doc.core_properties.subject = "Summer course release and instructional development proposal"
doc.core_properties.author = "Chase"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
