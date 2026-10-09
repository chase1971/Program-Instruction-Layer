from pathlib import Path
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(sys.argv[1])
OUTPUT = Path(sys.argv[2])


def set_text(paragraph, text, bold_prefix=None):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def replace_text_everywhere(doc, old, new):
    containers = [doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                containers.append(cell.paragraphs)
    for paragraphs in containers:
        for paragraph in paragraphs:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)


def paragraph_by_text(doc, text):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def paragraph_index_by_text(doc, text):
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == text:
            return index
    raise ValueError(f"Paragraph not found: {text}")


def remove_between(start_paragraph, end_paragraph):
    body = start_paragraph._p.getparent()
    current = start_paragraph._p.getnext()
    while current is not None and current is not end_paragraph._p:
        following = current.getnext()
        body.remove(current)
        current = following


def remove_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


doc = Document(SOURCE)

# Use the distinct Math Portal name throughout the proposal.
replace_text_everywhere(doc, "student portal", "Math Portal")
replace_text_everywhere(doc, "Student portal", "Math Portal")

# Make the consistency advantage over generic publisher and online resources explicit.
current_access_problem = paragraph_by_text(
    doc,
    "At present, students can search for help on their own, use a publisher system, ask an instructor, or consult materials distributed in an individual course. These options do not provide a consistent path. Students often do not know the name of the skill they need, which source matches the method used in class, or what they should do after watching an explanation. Instructors also lack one common destination they can recommend when a student needs targeted review.",
)
set_text(
    current_access_problem,
    "At present, students can search for help on their own, use a publisher system, ask an instructor, or consult materials distributed in an individual course. These options do not provide a consistent path. Students often do not know the name of the skill they need, which source matches the method used in class, or what they should do after watching an explanation. Publisher systems and online explanations may also use different notation, steps, or solution methods from those expected by the student's instructor. Instructors lack one common destination they can recommend when a student needs targeted review that is consistent with classroom instruction.",
)

# Reframe the existing platform around delivery plus evidence-based improvement.
delivery_heading = paragraph_by_text(doc, "Existing Delivery and Reporting Tools")
set_text(delivery_heading, "Existing Delivery Reporting and Data Collection Tools")

delivery_start = paragraph_index_by_text(doc, "Existing Delivery Reporting and Data Collection Tools")
delivery_paragraphs = doc.paragraphs[delivery_start + 1: delivery_start + 8]
delivery_updates = [
    "•  The Math Portal web application can deliver videos, reference pages, quizzes, guided activities, and tutorials.",
    "•  Student-specific access allows students to save work and resume an activity later.",
    "•  Selected activities provide immediate feedback and guided recovery after an error.",
    "•  Student-level and aggregate reporting can show attempt history, correct responses, common errors, help use, completion, and other patterns that identify where students are struggling.",
    "•  In-app surveys and other student feedback can help assess whether an activity is useful and whether its explanations, directions, or feedback are clear.",
    "•  Data collection and continuous improvement: Activity data and student feedback can be reviewed together to identify needed revisions, improve the app, and examine whether later changes produce clearer instruction, stronger completion, or better performance.",
    "•  A privacy-conscious design keeps student names out of the cloud activity database.",
]
for paragraph, text in zip(delivery_paragraphs, delivery_updates):
    prefix = "•  Data collection and continuous improvement: " if "Data collection and continuous improvement:" in text else None
    set_text(paragraph, text, bold_prefix=prefix)

# Hold preliminary evidence until the latest classroom results are ready.
evidence_heading = paragraph_by_text(doc, "Preliminary Student Evidence")
scope_heading = paragraph_by_text(doc, "Project Scope")
remove_between(evidence_heading, scope_heading)
blank = scope_heading.insert_paragraph_before("")
blank.paragraph_format.space_after = Pt(10)

# Make the purpose of app data explicit in the evaluation plan.
evaluation_heading = paragraph_by_text(doc, "Evaluation Plan")
evaluation_index = paragraph_index_by_text(doc, "Evaluation Plan")
evaluation_intro = doc.paragraphs[evaluation_index + 1]
set_text(
    evaluation_intro,
    "The initial evaluation will focus on whether students use the resources, whether they can complete the activities, where errors or requests for help occur, and whether students report that the materials improve their understanding and confidence. The available class sizes may not support strong causal claims, so the results will be used primarily for instructional improvement and feasibility assessment.",
)
evaluation_followup = doc.paragraphs[evaluation_index + 2]
set_text(
    evaluation_followup,
    "App-generated activity data and in-app student feedback will address two related questions: where are students struggling, and how should the instruction or app be improved? Patterns in responses, help use, completion, and student comments will guide revisions. Later use can then be reviewed to determine whether those changes improve clarity, completion, and student performance.",
)
evaluation_followup.paragraph_format.space_before = Pt(7)
evaluation_followup.paragraph_format.space_after = Pt(7)

comparison_paragraph = paragraph_by_text(
    doc,
    "Where a comparison is made between the Hub and another system, students should complete comparable work in both formats before being asked which they prefer. Any reported comparison will identify the task, sample size, exposure to each format, and limits of interpretation. This will avoid treating a preference question as evidence of effectiveness when the experiences were not equivalent.",
)
comparison_paragraph.paragraph_format.space_before = Pt(2)

# Add mobile access as a secondary future phase, not as a current capability or summer requirement.
requested_support = paragraph_by_text(doc, "Requested Support")
future_heading = requested_support.insert_paragraph_before("Future Math Portal Access")
future_heading.style = doc.styles["Heading 1"]
future_body_1 = requested_support.insert_paragraph_before(
    "A future phase could extend the College Algebra Skills Hub through the Math Portal, a mobile-friendly web application that students could access from a phone or other device. The Math Portal could provide convenient access to interactive practice, guided activities, tutorials, and selected reference materials."
)
future_body_1.style = doc.styles["Normal"]
future_body_2 = requested_support.insert_paragraph_before(
    "The Math Portal would complement rather than duplicate D2L. D2L would remain the complete home for the Hub, including its course organization, worksheets, notes, and longer video lessons. The Math Portal could provide direct access to app-based resources and link students to the corresponding D2L skill page when they need the full collection of materials. This future phase is secondary to the initial D2L development and is not included as a required summer deliverable."
)
future_body_2.style = doc.styles["Normal"]

customization_heading = requested_support.insert_paragraph_before("Future Instructor Specific Customization")
customization_heading.style = doc.styles["Heading 1"]
customization_body_1 = requested_support.insert_paragraph_before(
    "A later phase could add instructor-specific sections for skills where faculty members use different notation, problem-solving sequences, graphing methods, or required written steps. Students could follow a section-specific link or select their instructor so that the explanation, examples, and practice match the method used in their class. This would reduce the confusion that can occur when a publisher system or an online resource teaches a valid method that differs from the instructor's expectations."
)
customization_body_1.style = doc.styles["Normal"]
customization_body_2 = requested_support.insert_paragraph_before(
    "The shared College Algebra Skills Hub would remain the common foundation, and customized material would be added only when an instructional difference is important to student success. Developing these sections would require consultation with participating instructors, identification of the skills that genuinely need alternate treatments, faculty review of the resulting materials, and a plan for maintaining multiple versions. For those reasons, instructor-specific customization is a future enhancement rather than a required deliverable of the initial summer phase."
)
customization_body_2.style = doc.styles["Normal"]
customization_body_2.paragraph_format.keep_together = True

# Remove the risks section and its table in one structural operation.
risks_heading = paragraph_by_text(doc, "Risks and Responses")
approval_heading = paragraph_by_text(doc, "Approval Requested")
remove_between(risks_heading, approval_heading)
remove_element(risks_heading._p)

# Replace the long final checklist with the concise update reminder requested.
info_heading = paragraph_by_text(doc, "Information to Add Before Final Submission")
info_index = paragraph_index_by_text(doc, "Information to Add Before Final Submission")
info_body = doc.paragraphs[info_index + 1]
set_text(
    info_body,
    "Before submission, the proposal should be updated with the institution's required dates and format, the final funding period, the selected College Algebra skill priorities, and the latest classroom evidence.",
)
for paragraph in list(doc.paragraphs[info_index + 2:]):
    remove_element(paragraph._p)

# Keep repeating headers on tables that continue across pages.
for table in doc.tables:
    if table.rows:
        set_repeat_table_header(table.rows[0])

# Stabilize the manually numbered summer-work list across renderers.
for paragraph in doc.paragraphs:
    stripped = paragraph.text.strip()
    if any(stripped.startswith(f"{number}.") for number in range(1, 9)):
        paragraph.paragraph_format.line_spacing = 1.08
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.widow_control = True

long_item_six = paragraph_by_text(
    doc,
    "6.  Complete or refine selected full tutorials that can serve as models for later development, such as factoring or Gaussian elimination.",
)
set_text(
    long_item_six,
    "6.  Complete or refine selected full tutorials as models for later development.",
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
